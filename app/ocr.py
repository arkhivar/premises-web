import io
import re
import logging

log = logging.getLogger(__name__)


def detect_format(filename: str) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in ("pdf",):
        return "pdf"
    if ext in ("docx",):
        return "docx"
    if ext in ("jpg", "jpeg", "png", "gif", "webp", "bmp", "tiff", "tif"):
        return "image"
    return "unknown"


def extract_text(file_bytes: bytes, filename: str) -> str:
    fmt = detect_format(filename)
    if fmt == "pdf":
        return _pdf(file_bytes)
    if fmt == "docx":
        return _docx(file_bytes)
    if fmt == "image":
        return _ocr(file_bytes)
    return ""


def _pdf(file_bytes: bytes) -> str:
    import fitz

    doc = fitz.open(stream=file_bytes, filetype="pdf")
    text = ""
    try:
        for page in doc:
            text += page.get_text()
        if len(text.strip()) < 30:
            text = ""
            for page in doc:
                pix = page.get_pixmap(dpi=250)
                img_data = pix.tobytes("png")
                text += _ocr(img_data) + "\n"
    finally:
        doc.close()
    return text


def _docx(file_bytes: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(file_bytes))
    parts = []
    for p in doc.paragraphs:
        parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text for cell in row.cells)
            parts.append(row_text)
    return "\n".join(parts)


def _ocr(file_bytes: bytes) -> str:
    import pytesseract
    from PIL import Image

    img = Image.open(io.BytesIO(file_bytes))
    return pytesseract.image_to_string(img, lang="rus+eng")


# ── field extraction ────────────────────────────────────────


def parse_fields(raw_text: str) -> dict:
    result = {}
    t = _clean(raw_text)

    result.update(_extract_cadastral(t))
    result.update(_extract_area(t))
    result.update(_extract_floor(t))
    result.update(_extract_address(t, raw_text))
    result.update(_extract_type(t))
    result.update(_extract_name(t, result))

    return result


def _clean(text: str) -> str:
    t = text.replace("\r\n", "\n").replace("\r", "\n")
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t


# ── individual extractors ───────────────────────────────────

_CAD_PATTERNS = [
    re.compile(
        r"(?:кад(?:астровый|астров|\.)\s+(?:номер|№|н\.|знак)|КН)\s*[:]?\s*(\d{1,3}:\d{1,3}:\d{4,12}:\d{1,10})",
        re.I,
    ),
    re.compile(r"(\d{2}:\d{2}:\d{6,10}:\d{1,6})"),
]


def _extract_cadastral(text: str) -> dict:
    for pat in _CAD_PATTERNS:
        m = pat.search(text)
        if m:
            return {"cadastralNumber": m.group(1).strip()}
    return {}


_AREA_PATTERNS = [
    re.compile(
        r"площад[ьи]+\s*(?:помещения|объекта|здания|застройки)?\s*[:]?\s*"
        r"(\d+[.,]?\d*)\s*(?:кв[.]?\s*м[.]?|м[²2]|㎡|квм)?",
        re.I,
    ),
    re.compile(r"общ[ая]+\s*площад[ьи]*\s*[:]?\s*(\d+[.,]?\d*)", re.I),
]


def _extract_area(text: str) -> dict:
    for pat in _AREA_PATTERNS:
        m = pat.search(text)
        if m:
            area_str = m.group(1).replace(",", ".")
            try:
                return {"area": float(area_str)}
            except ValueError:
                pass
    return {}


_FLOOR_PATTERNS = [
    re.compile(r"этаж\s*(?:ность|№|номер)?\s*[:]?\s*(\d+)", re.I),
    re.compile(r"(?:расположен[ао]?\s+на|находится\s+на)\s+(\d+)\s*(?:-?м|ом)?\s*этаж", re.I),
]


def _extract_floor(text: str) -> dict:
    for pat in _FLOOR_PATTERNS:
        m = pat.search(text)
        if m:
            return {"floor": int(m.group(1))}
    return {}


_ADDR_PATTERNS = [
    re.compile(r"адрес\s*(?:\(местоположение\))?\s*[:]?\s*(.*?)(?:\n|Кадастровый|Площадь|Назначение|$)", re.I | re.S),
    re.compile(r"местоположение\s*[:]?\s*(.*?)(?:\n|Кадастровый|Площадь|$)", re.I | re.S),
    re.compile(r"местонахождение\s*[:]?\s*(.*?)(?:\n|Кадастровый|Площадь|$)", re.I | re.S),
]


def _extract_address(text: str, raw: str) -> dict:
    for pat in _ADDR_PATTERNS:
        m = pat.search(text)
        if m:
            addr = m.group(1).strip().rstrip(".,;: ")
            addr = re.sub(r"\s+", " ", addr)
            if len(addr) > 5:
                return {"address": addr}
    # Fallback: search for typical Russian address patterns
    city = re.search(r"г[.]\s*([А-ЯЁ][а-яё\s\-]+?)(?:[,]|\s+ул|\s+пр|\s+пер|\s+ш\.)", text)
    if city:
        idx = text.lower().find(city.group(0).lower())
        addr = text[idx:].split("\n")[0].strip().rstrip(".,;: ")
        if len(addr) > 5:
            return {"address": addr}
    return {}


_TYPE_MAP = [
    ("OFFICE", [r"офис", r"офисн", r"административн", r"контор", r"кабинет"]),
    ("APARTMENT", [r"квартир", r"апартамент", r"жилое\s*помещение", r"жилое\s*здание"]),
    ("COMMERCIAL", [r"коммерч", r"торгов", r"магазин", r"нежилое\s*помещение", r"нежилое\s*здание",
                    r"общепит", r"кафе", r"ресторан", r"гостиниц"]),
    ("WAREHOUSE", [r"склад", r"складск", r"ангар", r"бокс", r"гараж"]),
]


def _extract_type(text: str) -> dict:
    t = text.lower()
    for ptype, keywords in _TYPE_MAP:
        for kw in keywords:
            if re.search(kw, t):
                return {"type": ptype}
    # Default from assignment/naznachenie
    naz_match = re.search(r"назначение\s*[:]?\s*(.*?)(?:\n|$)", text, re.I)
    if naz_match:
        nazn = naz_match.group(1).strip().lower()
        for ptype, keywords in _TYPE_MAP:
            for kw in keywords:
                if re.search(kw, nazn):
                    return {"type": ptype}
    return {}


_NAME_PATTERNS = [
    re.compile(r"(?:наименование|название)\s*(?:объекта|помещения|здания)?\s*[:]?\s*(.*?)(?:\n|$)", re.I),
    re.compile(r"вид\s*объекта\s*(?:недвижимости)?\s*[:]?\s*(.*?)(?:\n|$)", re.I),
]


def _extract_name(text: str, existing: dict) -> dict:
    for pat in _NAME_PATTERNS:
        m = pat.search(text)
        if m:
            name = m.group(1).strip().rstrip(".,;: ")
            if len(name) > 1:
                return {"name": name}
    # Generate fallback name from type + address
    if "name" not in existing:
        type_names = {
            "OFFICE": "Офис",
            "APARTMENT": "Квартира",
            "COMMERCIAL": "Коммерческое помещение",
            "WAREHOUSE": "Склад",
        }
        prefix = type_names.get(existing.get("type", ""), "Помещение")
        addr = existing.get("address", "")
        if addr:
            return {"name": f"{prefix}, {addr[:60]}"}
    return {}


# ── tenant field extraction ─────────────────────────────────


def parse_tenant_fields(raw_text: str) -> dict:
    result = {}
    t = _clean(raw_text)

    result.update(_extract_tenant_type(t))
    result.update(_extract_tenant_name(t, result))
    result.update(_extract_inn(t))
    result.update(_extract_kpp(t))
    result.update(_extract_ogrn(t))
    result.update(_extract_legal_address(t))
    result.update(_extract_phone(t))
    result.update(_extract_email(t))
    result.update(_extract_bank_name(t))
    result.update(_extract_bik(t))
    result.update(_extract_bank_account(t))
    result.update(_extract_bank_corr_account(t))
    result.update(_extract_passport_series(t))
    result.update(_extract_passport_number(t))
    result.update(_extract_registration_address(t))

    return result


_TENANT_TYPE_PATTERNS = [
    (re.compile(r"\b(?:ООО|АО|ПАО|ЗАО|НАО|ОАО|НО)\b", re.I), "LEGAL_ENTITY"),
    (re.compile(r"\bИП\b", re.I), "ENTREPRENEUR"),
]


def _extract_tenant_type(text: str) -> dict:
    best_pos = len(text)
    best_type = None
    for pat, ttype in _TENANT_TYPE_PATTERNS:
        m = pat.search(text)
        if m and m.start() < best_pos:
            best_pos = m.start()
            best_type = ttype
    if best_type:
        return {"type": best_type}
    return {}


_TENANT_NAME_PATTERNS = [
    re.compile(
        r"(?:полное\s+)?наименование\s*(?:организации|компании|предприятия)?\s*[:]?\s*"
        r"(.*?)(?:\n|$)",
        re.I,
    ),
    re.compile(r"\bИП\s+([А-ЯЁ][а-яё]+(?:\s+[А-ЯЁ][а-яё]+){0,2})", re.I),
    re.compile(r"\b(?:ООО|АО|ПАО|ЗАО|НАО|ОАО|НО)\s+[\"«]?(.*?)[\"»]?(?:\s*[,.\n]|$)", re.I),
]


def _extract_tenant_name(text: str, existing: dict) -> dict:
    ttype = existing.get("type")
    patterns = list(_TENANT_NAME_PATTERNS)
    if ttype == "ENTREPRENEUR":
        patterns.sort(key=lambda p: 0 if p.pattern.startswith(r"\bИП") else 1)
    elif ttype == "LEGAL_ENTITY":
        patterns.sort(key=lambda p: 0 if r"ООО|АО" in p.pattern else 1)
    for pat in patterns:
        m = pat.search(text)
        if not m:
            continue
        name = m.group(1).strip().rstrip('.,;: ')
        if len(name) < 2:
            continue
        if ttype == "ENTREPRENEUR" and pat.pattern.startswith(r"\bИП"):
            return {"name": "ИП " + name}
        return {"name": name}
    return {}


_INN_PATTERN = re.compile(
    r"(?:ИИН|ИНН|inn)\s*[:\s]*\(?(\d[\d\s]{8,11}\d)", re.I
)


def _extract_inn(text: str) -> dict:
    m = _INN_PATTERN.search(text)
    if m:
        val = m.group(1).replace(" ", "").replace("-", "")
        if len(val) in (10, 12):
            return {"inn": val}
    fallback = re.search(r"\b(\d{10})\b", text)
    if fallback:
        return {"inn": fallback.group(1)}
    return {}


_KPP_PATTERN = re.compile(r"КПП\s*[:\s]*(\d[\d\s]{7,8}\d)", re.I)


def _extract_kpp(text: str) -> dict:
    m = _KPP_PATTERN.search(text)
    if m:
        val = m.group(1).replace(" ", "").replace("-", "")
        if len(val) == 9:
            return {"kpp": val}
    return {}


_OGRN_PATTERN = re.compile(
    r"ОГРНИП\s*[:\s]*(\d[\d\s]{13,14}\d)|"
    r"ОГРН\s*[:\s]*(\d[\d\s]{11,12}\d)",
    re.I,
)


def _extract_ogrn(text: str) -> dict:
    m = _OGRN_PATTERN.search(text)
    if m:
        val = (m.group(1) or m.group(2) or "").replace(" ", "").replace("-", "")
        if len(val) in (13, 15):
            return {"ogrn": val}
    return {}


_LEGAL_ADDR_PATTERNS = [
    re.compile(
        r"(?:юридическ(?:ий|ое|ая)\s+адрес|юр[.]?\s*адрес|адрес\s*(?:места\s+)?нахождени[яю])"
        r"\s*[:]?\s*(.*?)(?:\n|Фактический|Почтовый|ИНН|ОГРН|Тел|Телефон|$)",
        re.I | re.S,
    ),
    re.compile(
        r"(?:адрес\s*(?:места\s+)?нахождени[яю]|место\s+нахождени[яю])\s*[:]?\s*(.*?)(?:\n|ИНН|ОГРН|Тел|$)",
        re.I | re.S,
    ),
]


def _extract_legal_address(text: str) -> dict:
    for pat in _LEGAL_ADDR_PATTERNS:
        m = pat.search(text)
        if m:
            addr = m.group(1).strip().rstrip(".,;: ")
            addr = re.sub(r"\s+", " ", addr)
            if len(addr) > 5:
                return {"legalAddress": addr}
    return {}


_PHONE_PATTERN = re.compile(
    r"(?:тел(?:ефон)?|моб(?:ильный)?|факс)\s*[:\s]*"
    r"((?:\+7|8)(?:[\s\-()]*\d){10})",
    re.I,
)
_PHONE_FALLBACK = re.compile(r"(\+(?:7|3)\d[\s\-()]*\d{3}[\s\-()]*\d{3}[\s\-()]*\d{2}[\s\-()]*\d{2})\b")


def _extract_phone(text: str) -> dict:
    m = _PHONE_PATTERN.search(text)
    if m:
        return {"phone": m.group(1).strip()}
    m = _PHONE_FALLBACK.search(text)
    if m:
        return {"phone": m.group(1).strip()}
    return {}


_EMAIL_PATTERN = re.compile(
    r"(?:e-?mail|эл[.]?\s*почта|почта)\s*[:\s]*([^\s,;]+@[^\s,;]+)",
    re.I,
)
_EMAIL_FALLBACK = re.compile(r"\b([a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,})\b")


def _extract_email(text: str) -> dict:
    m = _EMAIL_PATTERN.search(text)
    if m:
        return {"email": m.group(1).strip()}
    m = _EMAIL_FALLBACK.search(text)
    if m:
        return {"email": m.group(1).strip()}
    return {}


_BANK_NAME_PATTERNS = [
    re.compile(r"(?:наименование\s+)?банк(?:а|е)?\s*[:]?\s*(.*?)(?:\n|БИК|К/с|Р/с|Корр|$)", re.I | re.S),
    re.compile(r"в\s+банке\s+(.*?)(?:\n|БИК|К/с|Р/с|Корр|$)", re.I | re.S),
]


def _extract_bank_name(text: str) -> dict:
    for pat in _BANK_NAME_PATTERNS:
        m = pat.search(text)
        if m:
            name = m.group(1).strip().rstrip('.,;: «»""')
            if len(name) > 3:
                return {"bankName": name}
    return {}


_BIK_PATTERN = re.compile(r"БИК\s*[:\s]*(\d{9})", re.I)


def _extract_bik(text: str) -> dict:
    m = _BIK_PATTERN.search(text)
    if m:
        return {"bankBik": m.group(1)}
    return {}


_BANK_ACCOUNT_PATTERN = re.compile(
    r"(?:р/с|расч[ёе]тн(?:ый|ая)\s+(?:сч[ёе]т|сч\.))\s*(?:№\s*)?[:\s]*"
    r"(\d[\d\s]{18,19}\d)",
    re.I,
)


def _extract_bank_account(text: str) -> dict:
    m = _BANK_ACCOUNT_PATTERN.search(text)
    if m:
        val = m.group(1).replace(" ", "").replace("-", "")
        if len(val) == 20:
            return {"bankAccount": val}
    return {}


_CORR_ACCOUNT_PATTERN = re.compile(
    r"(?:к/с|корр(?:еспондентск(?:ий|ая))?\s+(?:сч[ёе]т|сч\.))\s*(?:№\s*)?[:\s]*"
    r"(\d[\d\s]{18,19}\d)",
    re.I,
)


def _extract_bank_corr_account(text: str) -> dict:
    m = _CORR_ACCOUNT_PATTERN.search(text)
    if m:
        val = m.group(1).replace(" ", "").replace("-", "")
        if len(val) == 20:
            return {"bankCorrAccount": val}
    return {}


_PASSPORT_SERIES_PATTERN = re.compile(
    r"(?:серия)\s*(?:паспорта)?\s*[:\s]*(\d{2}\s*\d{2})", re.I
)


def _extract_passport_series(text: str) -> dict:
    m = _PASSPORT_SERIES_PATTERN.search(text)
    if m:
        return {"passportSeries": m.group(1).replace(" ", "")}
    return {}


_PASSPORT_NUMBER_PATTERN = re.compile(
    r"(?:номер)\s*(?:паспорта)?\s*[:\s]*(\d{6})", re.I
)


def _extract_passport_number(text: str) -> dict:
    m = _PASSPORT_NUMBER_PATTERN.search(text)
    if m:
        return {"passportNumber": m.group(1)}
    return {}


_REG_ADDR_PATTERNS = [
    re.compile(
        r"(?:адрес\s+регистрации|место\s+жительства|адрес\s+проживания|регистраци[яю]\s+по\s+месту\s+жительства)"
        r"\s*[:]?\s*(.*?)(?:\n|Паспорт|ИНН|Тел|Телефон|$)",
        re.I | re.S,
    ),
]


def _extract_registration_address(text: str) -> dict:
    for pat in _REG_ADDR_PATTERNS:
        m = pat.search(text)
        if m:
            addr = m.group(1).strip().rstrip(".,;: ")
            addr = re.sub(r"\s+", " ", addr)
            if len(addr) > 5:
                return {"registrationAddress": addr}
    return {}

